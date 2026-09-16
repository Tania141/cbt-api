"""Слой 2 → окончателния доклад.

РЕШЕНИЯ НА ОПЕРАТОРА (16.09.2026)
--------------------------------
1. В доклада влиза САМО потвърденото с хартията („✔ Сверено“ / „✎ Поправи“).
   Непотвърденото остава точки — докладът отива в ДНСК.
2. Паспортът (датите на съставяне, номерът на ЗК) е въведен от оператора и
   също се брои за потвърден. Когато паспортът и потвърденото в слой 2 дават
   РАЗЛИЧНИ стойности — вариант в): редът НЕ се попълва, докато не ги изравни,
   и системата казва защо. Никой от двата източника не печели мълчаливо.

Шаблонът не се пипа: в раздел Г 1.1 редовете са текст с точки, не полета.
Кодът намира параграфа по началото му и пише само след него.
"""
import copy, re
from datetime import date

from .sloy2 import n_data, n_nomer, _dd, _za_drug_stroezh

POTVARDENO = "потвърдено от оператора"


def _potv(f):
    """Стойността на факт, само ако е потвърдена с хартията."""
    if isinstance(f, dict) and f.get("sverka") == POTVARDENO and str(f.get("stoynost") or "").strip():
        return str(f["stoynost"]).strip()
    return None


def _data_potv(d):
    x = _potv(d.get("data"))
    return n_data(x) if x else None


# ── Редовете на Г 1.1 ─────────────────────────────────────────────────────────
# (етикет в доклада, vid_kod в слой 2, ключ в паспорта docDates)
REDOVE = [
    ("Протокол обр.2", "PROTOKOL2", "protokol2"),
    ("Протокол обр.3", "OBR3", "obrazec3"),
    ("Акт обр.14", "AKT14", "akt14"),
    ("Акт обр.15", "AKT15", "akt15"),
]


def _stoynosti(izvori):
    """{нормализирана стойност: [откъде]}"""
    g = {}
    for st, otkade in izvori:
        if st:
            g.setdefault(st, []).append(otkade)
    return g


def _reshi(etiket, izvori):
    """Една стойност → нея; няколко различни → None и причина (вариант в)."""
    g = _stoynosti(izvori)
    if not g:
        return None, None
    if len(g) == 1:
        return next(iter(g)), None
    razl = "; ".join(f"{k} ({', '.join(v)})" for k, v in g.items())
    return None, f"{etiket}: паспортът и потвърденото от документите се разминават — {razl}. Изравни ги."


def za_doklad(dokumenti, doc_dates=None, zk=None):
    """→ {"g11": {етикет: текст или None}, "spisatsi": {...}, "prichini": [...]}.

    `dokumenti` — прочетеното с приложени потвърждения (както го праща PWA);
    `doc_dates` — датите на съставяне от паспорта; `zk` — {number, date} от регистъра.
    """
    doc_dates, zk = doc_dates or {}, zk or {}
    sgrada = [d for d in dokumenti or [] if not _za_drug_stroezh(d)]
    g11, prichini = {}, []

    for etiket, kod, dd_kl in REDOVE:
        izvori = [(_dd(x), d.get("id") or d.get("fajl", "")) for d in sgrada
                  if d.get("vid_kod") == kod for x in [_data_potv(d)] if x]
        pasp = n_data(doc_dates.get(dd_kl))
        if pasp:
            izvori.append((_dd(pasp), "паспорт"))
        st, pr = _reshi(etiket, izvori)
        g11[etiket] = st
        if pr:
            prichini.append(pr)

    # Заповедна книга: номерът е на регистъра (само той раздава номера);
    # датата на заверка — паспорт срещу потвърдената дата на документ ЗК.
    izvori = [(_dd(x), d.get("id") or d.get("fajl", "")) for d in sgrada
              if d.get("vid_kod") == "ZK" for x in [_data_potv(d)] if x]
    pasp = n_data(doc_dates.get("zk_zaverka")) or n_data(zk.get("date"))
    if pasp:
        izvori.append((_dd(pasp), "паспорт"))
    zk_data, pr = _reshi("Заповедна книга — дата", izvori)
    if pr:
        prichini.append(pr)
    g11["Заповедна книга"] = {"nomer": str(zk.get("number") or "").strip() or None, "data": zk_data}

    # Актове обр. 7: брой и период. Слой 2 — само ако ВСИЧКИ прочетени актове
    # са с потвърдена дата; иначе периодът би бил от непроверени дати.
    akt7 = [d for d in sgrada if d.get("vid_kod") == "AKT7"]
    s2 = None
    if akt7:
        dati = [_data_potv(d) for d in akt7]
        if all(dati):
            s2 = (len(dati), _dd(min(dati)), _dd(max(dati)))
        else:
            prichini.append(f"Акт обр.7: {sum(not x for x in dati)} от {len(dati)} прочетени акта са без "
                            f"потвърдена дата — периодът не се попълва от документите.")
    p7 = doc_dates.get("akt7") or {}
    pasp7 = None
    if n_data(p7.get("ot")) and n_data(p7.get("do")):
        pasp7 = (int(p7["br"]) if str(p7.get("br", "")).isdigit() else None,
                 _dd(n_data(p7["ot"])), _dd(n_data(p7["do"])))
    if s2 and pasp7 and (s2[1:] != pasp7[1:] or (pasp7[0] is not None and pasp7[0] != s2[0])):
        prichini.append(f"Акт обр.7: паспортът дава {pasp7[0] or '?'} акта от {pasp7[1]} до {pasp7[2]}, "
                        f"потвърдените документи — {s2[0]} акта от {s2[1]} до {s2[2]}. Изравни ги.")
        g11["Акт обр.7"] = None
    else:
        g11["Акт обр.7"] = s2 or pasp7

    # ── Списъците: т. 3 и раздел В ──────────────────────────────────────────
    # Документ влиза, ако поне номерът или датата е потвърдена и няма
    # непотвърдена от двете.
    def red(d):
        n, dt = d.get("nomer"), d.get("data")
        imat = [f for f in (n, dt) if isinstance(f, dict) and str(f.get("stoynost") or "").strip()]
        if not imat or any(_potv(f) is None for f in imat):
            return None
        chasti = [str(d.get("vid") or "").strip() or "документ"]
        if _potv(n):
            chasti.append(f"№ {_potv(n)}")
        if _potv(dt):
            chasti.append(f"от {_dd(n_data(_potv(dt))) if n_data(_potv(dt)) else _potv(dt)} г.")
        izd = _potv(d.get("izdatel"))
        return " ".join(chasti) + (f", {izd}" if izd else "")

    spisatsi = {"laboratorii": [], "aktove": [], "stanovishta": [], "v_drugi": []}
    nepotv = 0
    for d in dokumenti or []:
        kod = d.get("vid_kod")
        # Декларациите за материали са в отделна таблица (deklaracii_tablica);
        # енергийният сертификат не е декларация за материал — отива в раздел В.
        kade = ("laboratorii" if kod == "IZMERVANE"
                else "aktove" if kod in ("AKT12",)
                else "v_drugi" if kod == "ENERGIEN"          # само в раздел В, не при мрежите
                else "stanovishta" if kod in ("STANOVISHTE", "PRISAEDINYAVANE") or _za_drug_stroezh(d)
                else None)
        if not kade:
            continue
        t = red(d)
        if t:
            if t not in spisatsi[kade]:
                spisatsi[kade].append(t)
        else:
            nepotv += 1
    if nepotv:
        prichini.append(f"{nepotv} прочетени документа за т. 3 и раздел В не са влезли — "
                        f"номерът или датата им не е сверена с хартията.")
    deklaracii, bel = deklaracii_tablica(dokumenti)
    prichini += bel
    sobstvenici, sobstvenost, do, bel = sobstvenost_ot_aktove(dokumenti)
    prichini += bel
    return {"g11": g11, "spisatsi": spisatsi, "deklaracii": deklaracii, "sobstvenici": sobstvenici,
            "sobstvenost": sobstvenost, "do": do, "prichini": prichini}


# ── Собствеността: таблицата, списъкът, „ДО:“ ────────────────────────────────
# Решение на оператора (16.09.2026): таблицата „обект → собственик → нот. акт“
# е „най-трудоемка и за описване, и за прочитане“ — сглобява се от сверените
# нотариални актове. Влиза само свереното; частично сверено — с точки.

KOLONI_SOBSTVENICI = ("Обект", "Собственик", "Нот. акт, вписване по ЗС/ПВ, Служба по вписванията — гр. София")


def _vpisvane(d):
    """Кратко за колоната: вписването, а без него — самият нотариален акт."""
    ch = []
    for kl, et in (("vp_akt", "акт № "), ("vp_tom", "том "), ("vp_delo", "дело № ")):
        if _potv(d.get(kl)):
            ch.append(et + _potv(d.get(kl)))
    if not ch:
        for kl, et in (("nomer", "нот. акт № "), ("nt_tom", "том "), ("nt_delo", "дело № ")):
            if _potv(d.get(kl)):
                ch.append(et + _potv(d.get(kl)))
    return ", ".join(ch)


def _pal_zapis(d):
    """Пълният запис за „Документи за собственост“ — само от сверени части."""
    if not _potv(d.get("nomer")):
        return None
    t = f"Нотариален акт № {_potv(d.get('nomer'))}"
    for kl, et in (("nt_tom", ", том "), ("nt_reg", ", рег. № "), ("nt_delo", ", дело № ")):
        if _potv(d.get(kl)):
            t += et + _potv(d.get(kl))
    if _potv(d.get("data")):
        t += f" от {_potv(d.get('data'))} г."
    if _potv(d.get("nt_notarius")):
        t += f" на нотариус {_potv(d.get('nt_notarius'))}"
    vp = [et + _potv(d.get(kl)) for kl, et in (("vp_vh_reg", "вх. рег. № "), ("vp_akt", "акт "),
                                               ("vp_tom", "том "), ("vp_delo", "дело № ")) if _potv(d.get(kl))]
    if vp:
        t += ", вписан в СВ с " + ", ".join(vp)
    return t


def sobstvenost_ot_aktove(dokumenti):
    """→ (редове на таблицата, пълни записи, собственици за „ДО:“, бележки)"""
    redove, zapisi, imena, belezhki = [], [], [], []
    nesvereni_redove = nesvereni_aktove = 0
    for d in dokumenti or []:
        if d.get("vid_kod") != "NOT_AKT":
            continue
        z = _pal_zapis(d)
        if z:
            if z not in zapisi:
                zapisi.append(z)
        else:
            nesvereni_aktove += 1
        akt = _vpisvane(d) or "…………"
        for r in d.get("razpredelenie") or []:
            ob, sob = _potv(r.get("obekti")), _potv(r.get("sobstvenici"))
            if not ob or not sob:
                nesvereni_redove += 1
                continue
            redove.append({"obekt": ob, "sobstvenik": sob, "akt": akt})
            for ime in re.split(r"\s*;\s*", sob):
                if ime.strip() and ime.strip() not in imena:
                    imena.append(ime.strip())
    if nesvereni_redove:
        belezhki.append(f"Собственици: {nesvereni_redove} реда не са влезли в таблицата — обектите или "
                        f"собствениците не са сверени с хартията.")
    if nesvereni_aktove:
        belezhki.append(f"Документи за собственост: {nesvereni_aktove} нотариални акта без сверен номер — не са влезли.")
    return redove, zapisi, ", ".join(imena), belezhki


# ── Таблицата на декларациите ────────────────────────────────────────────────
# Решение на оператора (16.09.2026): ЕДНА таблица в „Декларации за
# съответствие на вложените материали“ — Материал/изделие · Производител/
# Доставчик · Вид (с №) · Дата на издаване. Шест доставки бетон по една
# декларация са един ред с шест дати. Различно изписване не се слива — пита.

KOLONI_DEKLARACII = ("Материал/изделие", "Производител/Доставчик",
                     "Вид (Сертификат/Декларация/друго)", "Дата на издаване")


def _kl(s):
    return re.sub(r"[\W_]+", " ", str(s or "").lower()).strip()


def deklaracii_tablica(dokumenti):
    """Сверените декларации → (редове, бележки). Ред = материал+производител+№."""
    grupi, nevlezli, belezhki = {}, 0, []
    for d in dokumenti or []:
        if d.get("vid_kod") != "DEKLARACIA":
            continue
        mat, dt_raw = _potv(d.get("material")), _potv(d.get("data"))
        dt = n_data(dt_raw) if dt_raw else None
        if not mat or not dt:
            nevlezli += 1
            continue
        proizv = _potv(d.get("proizvoditel")) or _potv(d.get("izdatel")) or ""
        nomer = _potv(d.get("nomer")) or ""
        vid = str(d.get("vid") or "Декларация").strip()
        vid_txt = f"{vid} № {nomer}" if nomer and nomer not in vid else vid
        kl = (_kl(mat), _kl(proizv), n_nomer(nomer))
        red = grupi.setdefault(kl, {"material": mat, "proizvoditel": proizv, "vid": vid_txt, "dati": set()})
        red["dati"].add(dt)
    po_nomer = {}
    for (_, _, n), red in grupi.items():
        if n:
            po_nomer.setdefault(n, []).append(red)
    for n, rs in po_nomer.items():
        if len(rs) > 1:
            belezhki.append(f"Декларация № {n}: изписана различно — "
                            + " / ".join(f"„{x['material']}“, {x['proizvoditel'] or '—'}" for x in rs)
                            + ". Не са слети — провери дали е един и същ материал.")
    if nevlezli:
        belezhki.append(f"Декларации: {nevlezli} не са влезли в таблицата — материалът или датата "
                        f"не е сверена с хартията.")
    redove = [{**r, "dati": ", ".join(_dd(x) for x in sorted(r["dati"])) + " г."} for r in grupi.values()]
    redove.sort(key=lambda r: (r["material"].lower(), r["dati"]))
    return redove, belezhki


def vstavi_deklaracii(doc, redove, marker="{{ОДАИ_Декларации}}"):
    return vstavi_tablica(doc, marker, KOLONI_DEKLARACII, ("material", "proizvoditel", "vid", "dati"),
                          redove, (2600, 2300, 3200, 1765))


def vstavi_sobstvenici(doc, redove, marker="{{ОДАИ_Собственици}}"):
    return vstavi_tablica(doc, marker, KOLONI_SOBSTVENICI, ("obekt", "sobstvenik", "akt"),
                          redove, (2800, 3300, 3765))


def vstavi_tablica(doc, marker, koloni, klyuchove, redove, shirini):
    """Таблица на мястото на маркера. Без редове — маркерът остава (излиза с точки)."""
    if not redove:
        return False
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips
    par = next((p for p in doc.paragraphs if marker in p.text), None)
    if par is None:
        return False
    # ширините са в dxa, общо 9865 — колкото главата
    tbl = doc.add_table(rows=1 + len(redove), cols=len(koloni))
    t = tbl._tbl
    tblPr = t.tblPr
    granici = OxmlElement("w:tblBorders")
    for strana in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{strana}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
            b.set(qn(f"w:{k}"), v)
        granici.append(b)
    tblPr.append(granici)
    stil = par.style
    for i, red in enumerate([dict(zip(klyuchove, koloni))] + redove):
        for j, kl in enumerate(klyuchove):
            kl_ = tbl.rows[i].cells[j]
            kl_.width = Twips(shirini[j])
            p = kl_.paragraphs[0]
            p.style = stil
            run = p.add_run(str(red.get(kl) or "…………"))
            if i == 0:
                run.bold = True
    if tbl.rows:
        # заглавният ред се повтаря на всяка страница
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trPr.append(h)
    par._p.addprevious(t)
    par._p.getparent().remove(par._p)
    return True


# ── Блоковете на шаблона „ОД с АИ“ (по ПЕТРАКИЕВ) ────────────────────────────

def blokove_od_ai(rez, d):
    """Резултатът от za_doklad + паспортът → стойностите на {{ОДАИ_…}}.

    Празното остава с точките от build_placeholders — нищо не се измисля.
    Собствениците, договорите и показателите идват в следващите стъпки.
    """
    g, sp = rez["g11"], rez["spisatsi"]
    red = lambda xs: "\n".join(f"- {x}" for x in xs if x)
    out = {}

    zakon = []
    if d.get("РС_Номер") or d.get("РС_Дата"):
        rs = f"Разрешение за строеж № {d.get('РС_Номер', '').strip()} от {d.get('РС_Дата', '').strip()} г."
        if d.get("РС_Издател"):
            rs += f", издадено от {d['РС_Издател'].strip()}"
        if d.get("РС_ВСила"):
            rs += f", влязло в сила на {d['РС_ВСила'].strip()} г."
        zakon.append(rs)
    if g.get("Протокол обр.2"):
        zakon.append(f"Протокол за откриване на строителна площадка и за определяне на строителна линия "
                     f"и ниво от {g['Протокол обр.2']} г.")
    if g.get("Протокол обр.3"):
        zakon.append(f"Констативен акт обр. 3 за установяване съответствието на строежа с издадените строителни "
                     f"книжа и за това, че подробният устройствен план е приложен по отношение на застрояването "
                     f"от {g['Протокол обр.3']} г.")
    z = g.get("Заповедна книга") or {}
    if z.get("nomer") or z.get("data"):
        zakon.append("Заповедна книга" + (f" № {z['nomer']}" if z.get("nomer") else "")
                     + (f" от {z['data']} г." if z.get("data") else ""))
    if zakon:
        out["{{ОДАИ_Законосъобразно}}"] = red(zakon)

    aktove = []
    a7 = g.get("Акт обр.7")
    if a7:
        aktove.append(f"Актове обр. 7 за приемане на конструкцията по нива и елементи"
                      + (f" — {a7[0]} бр." if a7[0] else "") + f" от {a7[1]} г. до {a7[2]} г.")
    if g.get("Акт обр.14"):
        aktove.append(f"Акт обр. 14 за приемане на строителната конструкция от {g['Акт обр.14']} г.")
    if g.get("Акт обр.15"):
        aktove.append(f"Акт обр. 15 от {g['Акт обр.15']} г.")
    aktove += sp.get("aktove") or []
    if aktove:
        out["{{ОДАИ_Актове}}"] = red(aktove)

    if sp.get("stanovishta"):
        out["{{ОДАИ_Мрежи}}"] = red(sp["stanovishta"])
    if sp.get("stanovishta") or sp.get("v_drugi"):
        out["{{ОДАИ_В}}"] = red((sp.get("stanovishta") or []) + (sp.get("v_drugi") or []))
    if sp.get("laboratorii"):
        out["{{ОДАИ_Изпитвания}}"] = red(sp["laboratorii"])
    if rez.get("sobstvenost"):
        out["{{ОДАИ_Собственост}}"] = red(rez["sobstvenost"])
    if rez.get("do"):
        out["{{ОДАИ_ДО}}"] = rez["do"]
    return out


# ── Техническото описание на строежа ─────────────────────────────────────────
# Решение на оператора (16.09.2026): Акт 15 се мени, докато се пише ОД, но
# описанието на строежа е едно и също — само документите се дописват. Затова не
# се копира от Акт 15 в ОД, а се пази ВЕДНЪЖ в обекта и оттам го взимат двата
# (засега — ОД с АИ). Тук: как се изрязва от готов Акт 15 и как се слага в ОД.

NACHALO, KRAI = "===ТЕХНИЧЕСКО_НАЧАЛО===", "===ТЕХНИЧЕСКО_КРАЙ==="


def izrezhi_tehnichesko(redove):
    """Редовете на Акт 15 → текстът на описанието, дословно.

    Между маркерите, ако ги има (Акт 15 от шаблона); иначе от „Строежът
    представлява“ до „Въз основа на горните констатации“. None, ако не се намира.
    """
    redove = [str(r).rstrip() for r in redove]
    ch = lambda r: _nachalo(r)
    i = next((k for k, r in enumerate(redove) if ch(r) == NACHALO), None)
    j = next((k for k, r in enumerate(redove) if ch(r) == KRAI), None)
    if i is None or j is None or j <= i:
        # Истинският Акт 15 на оператора (ДЖИХАТ): „1. По издадените строителни
        # книжа…“ са документите (менят се); „2. По изпълнението на строежа:“ до
        # „Въз основа…“ е описанието (не се мени). Шаблонът — „Строежът представлява“.
        nach = re.compile(r"^(\d+\.\s*)?По изпълнението на строежа|^Строежът представлява")
        i = next((k for k, r in enumerate(redove) if nach.match(ch(r))), None)
        j = next((k for k, r in enumerate(redove) if ch(r).startswith("Въз основа на горните констатации")), None)
        if i is None or j is None or j <= i:
            return None
        izrez = redove[i:j]
        # „2. По изпълнението на строежа:“ е заглавие — в ОД си има свое („Б. По
        # изпълнение на СМР:“), да не се повтаря.
        if re.match(r"^(\d+\.\s*)?По изпълнението на строежа:?$", ch(izrez[0])):
            izrez = izrez[1:]
    else:
        izrez = redove[i + 1:j]
    # празни редове в началото и края — без тях
    while izrez and not izrez[0].strip():
        izrez.pop(0)
    while izrez and not izrez[-1].strip():
        izrez.pop()
    return "\n".join(izrez) or None


def redove_ot_fajl(ime, raw):
    """.docx или PDF с текст → редове. Сканиран PDF → None (няма текст)."""
    ext = ime.lower().rsplit(".", 1)[-1] if "." in ime else ""
    import io
    if ext == "docx":
        from docx import Document
        return [p.text for p in Document(io.BytesIO(raw)).paragraphs]
    if ext == "pdf":
        import pymupdf
        d = pymupdf.open(stream=raw, filetype="pdf")
        tekst = "\n".join(d[i].get_text() for i in range(d.page_count))
        return tekst.split("\n") if len(tekst.strip()) > 40 * max(d.page_count, 1) else None
    raise ValueError("дай Акт 15 като .docx или PDF с текст (стар .doc — запиши го като .docx)")


def vmukni_tehnichesko(doc, tekst):
    """Слага описанието между маркерите в ОД; маркерите и заместителят изчезват.
    Без текст — шаблонът остава както е."""
    if not (tekst or "").strip():
        return False
    from docx.text.paragraph import Paragraph
    pars = list(doc.paragraphs)
    i = next((k for k, p in enumerate(pars) if _nachalo(p.text) == NACHALO), None)
    j = next((k for k, p in enumerate(pars) if _nachalo(p.text) == KRAI), None)
    if i is None or j is None or j <= i:
        return False
    # Между маркерите има и истинско съдържание — в ОД заглавието „Б. По
    # изпълнение на СМР:“ стои вътре. Сменя се САМО заместителят „[ … ]“;
    # маркерите се махат, всичко друго остава.
    zam = next((pars[k] for k in range(i + 1, j) if _nachalo(pars[k].text).startswith("[")), None)
    if zam is None:
        return False
    redove = tekst.split("\n")
    _pishi(zam, redove[0])
    posleden = zam
    for r in redove[1:]:
        posleden = _sled(posleden, r)
    for k in (j, i):
        pars[k]._p.getparent().remove(pars[k]._p)
    return True


# ── Записване в .docx ─────────────────────────────────────────────────────────

def _pishi(par, tekst):
    """Сменя текста на параграф, като пази оформлението на първия run."""
    runs = par.runs
    if not runs:
        par.add_run(tekst)
        return
    runs[0].text = tekst
    for r in runs[1:]:
        r.text = ""


def _sled(par, tekst):
    """Нов параграф след `par`, със същото оформление."""
    novo = copy.deepcopy(par._p)
    par._p.addnext(novo)
    from docx.text.paragraph import Paragraph
    p = Paragraph(novo, par._parent)
    _pishi(p, tekst)
    return p


def _nachalo(t):
    return re.sub(r"\s+", " ", t or "").strip()


def zapishi(doc, rez):
    """Попълва Г 1.1 и списъците. Празното остава с точките на шаблона."""
    g11, sp = rez["g11"], rez["spisatsi"]
    tochki = "................................."
    for par in list(doc.paragraphs):
        t = _nachalo(par.text)
        for etiket in ("Протокол обр.2", "Протокол обр.3", "Акт обр.14", "Акт обр.15"):
            if t.startswith(f"{etiket} от"):
                if g11.get(etiket):
                    _pishi(par, f"{etiket} от {g11[etiket]} г.")
        if t.startswith("Заповедна книга №"):
            z = g11.get("Заповедна книга") or {}
            if z.get("nomer") or z.get("data"):
                _pishi(par, f"Заповедна книга № {z.get('nomer') or '..................'} "
                            f"от {(z['data'] + ' г.') if z.get('data') else tochki}")
        if t.startswith("Акт обр.7 от"):
            a7 = g11.get("Акт обр.7")
            if a7:
                br = f"{a7[0]} бр. " if a7[0] else ""
                _pishi(par, f"Актове обр.7 — {br}от {a7[1]} г. до {a7[2]} г.")
        for zagl, kl in (("Протоколи от Акредитирани лаборатории:", "laboratorii"),
                         ("Декларации за екслоатационни характеристики", "deklaracii"),
                         ("Актове и протоколи по време на строителството:", "aktove")):
            if t.startswith(zagl) and sp.get(kl):
                posleden = par
                for x in sp[kl]:
                    posleden = _sled(posleden, f"– {x}")
        if t.startswith("При извършеното строителство са спазени всички нормативни изисквания") and sp.get("stanovishta"):
            # следващият параграф е „.......“ — там отиват становищата
            sledvasht = par._p.getnext()
            from docx.text.paragraph import Paragraph
            if sledvasht is not None and _nachalo(Paragraph(sledvasht, par._parent).text).strip(".") == "":
                p0 = Paragraph(sledvasht, par._parent)
                _pishi(p0, f"– {sp['stanovishta'][0]}")
                posleden = p0
                for x in sp["stanovishta"][1:]:
                    posleden = _sled(posleden, f"– {x}")
    return doc
